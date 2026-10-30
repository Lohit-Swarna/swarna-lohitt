"""Source-partition replacement, row accountability and exact-money controls."""

from __future__ import annotations

import csv
import hashlib
import json
import sqlite3
from collections import Counter
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Iterable, Iterator

SCALE = 10_000
MAX_INTEGER = 2**63 - 1
MAX_PRICE = Decimal("922337203685477.5807")
REQUIRED = ("Invoice", "StockCode", "Description", "Quantity", "InvoiceDate", "Price", "Customer ID", "Country")
SQL_DIR = Path(__file__).resolve().parents[2] / "sql"


def fingerprint(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def identifier(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    try:
        number = Decimal(text)
        if number.is_finite() and number.copy_abs() <= MAX_INTEGER and number == number.to_integral_value():
            return str(int(number))
    except InvalidOperation:
        pass
    return text


def normalize(row: dict) -> dict:
    if set(row) != set(REQUIRED):
        raise ValueError("unexpected_field_count")
    invoice, product = identifier(row["Invoice"]), identifier(row["StockCode"])
    if not invoice or not product:
        raise ValueError("missing_invoice_or_product")
    try:
        quantity = Decimal(str(row["Quantity"]))
        price = Decimal(str(row["Price"]))
    except (InvalidOperation, ValueError):
        raise ValueError("invalid_numeric") from None
    if not quantity.is_finite() or not price.is_finite():
        raise ValueError("non_finite_numeric")
    if quantity != quantity.to_integral_value() or quantity == 0:
        raise ValueError("quantity_not_nonzero_integer")
    if price < 0:
        raise ValueError("negative_price")
    if quantity.copy_abs() > MAX_INTEGER or price > MAX_PRICE:
        raise ValueError("numeric_overflow")
    # Decimal multiplication honors the current context and can silently round long
    # input before a precision comparison. Derive fixed-point units from the exact
    # coefficient/exponent instead; insignificant trailing zeros are permitted.
    if price.is_zero():
        unit_value = 0
    else:
        _, digits, exponent = price.as_tuple()
        coefficient = "".join(str(digit) for digit in digits)
        trimmed = coefficient.rstrip("0")
        exponent += len(coefficient) - len(trimmed)
        if exponent < -4:
            raise ValueError("price_precision_over_four_decimals")
        if len(trimmed) + exponent + 4 > 19:
            raise ValueError("numeric_overflow")
        unit_value = int(trimmed) * 10 ** (exponent + 4)
    q = int(quantity)
    if abs(q * unit_value) > MAX_INTEGER:
        raise ValueError("numeric_overflow")
    date = row["InvoiceDate"]
    if isinstance(date, datetime):
        timestamp = date
    else:
        try:
            timestamp = datetime.fromisoformat(str(date))
        except (TypeError, ValueError):
            raise ValueError("invalid_timestamp") from None
    # Original data has naive local timestamps; never invent UTC conversion.
    if timestamp.tzinfo is not None:
        raise ValueError("timezone_aware_timestamp_not_in_contract")
    country = str(row["Country"] or "").strip()
    if not country:
        raise ValueError("missing_country")
    return {
        "invoice": invoice,
        "product": product,
        "description": str(row["Description"] or "").strip(),
        "quantity": q,
        "unit_value": unit_value,
        "line_value": q * unit_value,
        "timestamp": timestamp.isoformat(timespec="seconds"),
        "date": timestamp.date().isoformat(),
        "customer": identifier(row["Customer ID"]) or "UNKNOWN",
        "country": country,
        "is_cancellation": int(invoice.upper().startswith("C")),
    }


def read_csv(path: Path) -> Iterator[tuple[int, dict]]:
    with path.open(encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        header = reader.fieldnames or []
        if len(header) != len(REQUIRED) or len(set(header)) != len(REQUIRED) or set(header) != set(REQUIRED):
            raise ValueError("CSV schema differs from the eight-field source contract")
        for index, row in enumerate(reader, start=2):
            yield index, row


def read_xlsx(path: Path, sheet: str, limit: int | None = None) -> Iterator[tuple[int, dict]]:
    import openpyxl  # Optional, deliberately absent from the offline fixture path.

    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in book.sheetnames:
            raise ValueError("Unknown worksheet")
        rows = book[sheet].iter_rows(values_only=True)
        header = tuple(next(rows))
        if set(header) != set(REQUIRED) or len(header) != len(REQUIRED):
            raise ValueError("XLSX schema differs from source contract")
        for index, values in enumerate(rows, start=2):
            if limit is not None and index - 2 >= limit:
                break
            yield index, dict(zip(header, values))
    finally:
        book.close()


def connect(path: Path) -> sqlite3.Connection:
    path.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(path)
    connection.execute("PRAGMA foreign_keys=ON")
    connection.executescript((SQL_DIR / "schema.sql").read_text(encoding="utf-8"))
    return connection


def load(connection: sqlite3.Connection, records: Iterable[tuple[int, dict]], *, source_key: str, sha256: str, scope: str) -> dict:
    if not source_key or len(sha256) != 64:
        raise ValueError("Source key and SHA-256 are required")
    # BEGIN/commit includes deletes, writes and manifest: failed rebuild preserves prior facts.
    connection.execute("BEGIN")
    accepted = rejected = 0
    seen_ordinals = set()
    try:
        connection.execute("DELETE FROM fact_line WHERE source_key=?", (source_key,))
        connection.execute("DELETE FROM quarantine WHERE source_key=?", (source_key,))
        for row_number, row in records:
            if isinstance(row_number, bool) or not isinstance(row_number, int) or not 0 < row_number <= MAX_INTEGER:
                raise ValueError("source_row_must_be_positive_integer")
            if row_number in seen_ordinals:
                raise ValueError("duplicate_source_row_identity")
            seen_ordinals.add(row_number)
            try:
                item = normalize(row)
            except (ValueError, KeyError) as error:
                reason = str(error) if isinstance(error, ValueError) else "missing_required_field"
                connection.execute("INSERT INTO quarantine VALUES(?,?,?)", (source_key, row_number, reason))
                rejected += 1
                continue
            connection.execute("INSERT OR IGNORE INTO dim_date VALUES(?,?,?,?)", (item["date"], int(item["date"][:4]), item["date"][:7], int(item["date"][8:10])))
            connection.execute("INSERT INTO dim_product VALUES(?,?) ON CONFLICT(product_code) DO UPDATE SET description=excluded.description", (item["product"], item["description"]))
            connection.execute("INSERT OR IGNORE INTO dim_customer VALUES(?,?)", (item["customer"], int(item["customer"] == "UNKNOWN")))
            connection.execute("INSERT OR IGNORE INTO dim_country VALUES(?)", (item["country"],))
            connection.execute("INSERT INTO fact_line VALUES(?,?,?,?,?,?,?,?,?,?,?,?)", (source_key, row_number, item["invoice"], item["product"], item["customer"], item["country"], item["date"], item["timestamp"], item["quantity"], item["unit_value"], item["line_value"], item["is_cancellation"]))
            accepted += 1
        connection.execute("INSERT INTO source_manifest VALUES(?,?,?,?,?,?) ON CONFLICT(source_key) DO UPDATE SET sha256=excluded.sha256,scope=excluded.scope,source_rows=excluded.source_rows,accepted_rows=excluded.accepted_rows,rejected_rows=excluded.rejected_rows", (source_key, sha256, scope, accepted + rejected, accepted, rejected))
        result = report(connection, source_key)
        connection.commit()
    except BaseException:
        connection.rollback()
        raise
    return result


def pounds(units: int) -> str:
    whole, fractional = divmod(abs(units), SCALE)
    return f"{'-' if units < 0 else ''}{whole}.{fractional:04d}"


def report(connection: sqlite3.Connection, source_key: str) -> dict:
    connection.row_factory = sqlite3.Row
    manifest = connection.execute("SELECT * FROM source_manifest WHERE source_key=?", (source_key,)).fetchone()
    if manifest is None:
        raise ValueError("Source partition has not been loaded")
    controls = dict(connection.execute((SQL_DIR / "controls.sql").read_text(encoding="utf-8"), {"source_key": source_key}).fetchone())
    reasons = dict(connection.execute("SELECT reason,COUNT(*) FROM quarantine WHERE source_key=? GROUP BY reason ORDER BY reason", (source_key,)).fetchall())
    monthly = [dict(r) for r in connection.execute((SQL_DIR / "monthly.sql").read_text(encoding="utf-8"), {"source_key": source_key})]
    for row in monthly:
        row["signed_line_value_gbp"] = pounds(row.pop("signed_line_value_units"))
    if controls["accepted_rows"] != manifest["accepted_rows"]:
        raise RuntimeError("Accepted-row control failed")
    if manifest["source_rows"] != controls["accepted_rows"] + sum(reasons.values()):
        raise RuntimeError("Source-row reconciliation failed")
    if controls["signed_line_value_units"] != controls["positive_line_value_units"] + controls["negative_line_value_units"]:
        raise RuntimeError("Signed-value reconciliation failed")
    if connection.execute("PRAGMA foreign_key_check").fetchall():
        raise RuntimeError("Foreign-key validation failed")
    return {"schema_version": 1, "currency": "GBP", "money_scale": SCALE, "source": dict(manifest), "controls": controls, "quarantine_reasons": reasons, "monthly": monthly}


def write_report(result: dict, directory: Path) -> None:
    directory.mkdir(parents=True, exist_ok=True)
    (directory / "run_manifest.json").write_text(json.dumps(result, indent=2) + "\n", encoding="utf-8", newline="\n")
    controls, source = result["controls"], result["source"]
    lines = ["# Governed BI: measured data-foundation report", "", "This is a retrospective source-processing demonstration, not a production deployment or business-impact claim.", "", f"**Scope:** {source['scope']}", "", f"**Source key:** `{source['source_key']}`", "", f"**Input SHA-256:** `{source['sha256']}`", "", "| Control | Measured value |", "|---|---:|", f"| Rows inspected | {source['source_rows']:,} |", f"| Accepted lines | {source['accepted_rows']:,} |", f"| Quarantined lines | {source['rejected_rows']:,} |", f"| Known-customer lines | {controls['known_customer_lines']:,} |", f"| Missing-customer lines retained | {controls['unknown_customer_lines']:,} |", f"| Cancellation-invoice lines | {controls['cancellation_invoice_lines']:,} |", f"| Negative-quantity lines | {controls['negative_quantity_lines']:,} |", f"| Signed line value, GBP | {pounds(controls['signed_line_value_units'])} |", f"| Positive line value, GBP | {pounds(controls['positive_line_value_units'])} |", f"| Negative line value, GBP | {pounds(controls['negative_line_value_units'])} |", f"| Source date range | {controls['first_date']} to {controls['last_date']} |", "", "Signed line value is quantity × stated unit price, preserving the source sign. It is not audited recognized revenue, profitability or savings. Cancellation invoices are flagged separately; unmatched reversals are not inferred. Figures retain four decimal places and are not rounded per line.", "", "## Quarantine accounting", "", "| Reason | Lines |", "|---|---:|"]
    lines += [f"| {reason} | {count:,} |" for reason, count in result["quarantine_reasons"].items()]
    if not result["quarantine_reasons"]:
        lines.append("| None | 0 |")
    lines += ["", "## Monthly source activity", "", "| Month | Lines | Signed line value, GBP |", "|---|---:|---:|"]
    lines += [f"| {r['month']} | {r['line_count']:,} | {r['signed_line_value_gbp']} |" for r in result["monthly"]]
    lines += ["", "## Decisions supported and limits", "", "An analyst can reconcile accepted and quarantined rows, separate missing-customer coverage from all-line totals, and investigate cancellation/negative-quantity activity before publishing a measure. A bounded sequential slice is not representative of the full business: no customer concentration or retention recommendation is inferred from this slice. Customer identifiers and individual rows are deliberately excluded from this report.", "", "See the source provenance, metric contract and reproduction instructions in the project README. Native Power BI report authoring, DAX execution, stakeholder validation and production operation remain uncompleted milestones."]
    (directory / "EXECUTIVE_REPORT.md").write_text("\n".join(lines) + "\n", encoding="utf-8", newline="\n")
